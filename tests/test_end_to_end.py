"""End-to-end regression: sample artifacts → bundle → HTML → assertions."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from qa_dashboard import SCHEMA_VERSION
from qa_dashboard._paths import make_test_id, ms_between_iso, sanitize_nodeid
from qa_dashboard.bundle.schemas import (
    BundleValidationError,
    validate_manifest,
    validate_test_record,
)
from qa_dashboard.report import export as export_mod
from qa_dashboard.report.generator import (
    generate_index_html,
    generate_standalone_test_html,
)


@pytest.fixture
def sample_run(tmp_path):
    """Generate a small but complete sample run in tmp_path."""
    from tests.fixtures.make_sample_run import make_sample

    run_dir = make_sample(
        out_root=tmp_path,
        num_tests=8,
        fail_rate=0.25,
        steps_per_test=5,
        seed=123,
    )
    return run_dir


def test_sanitize_nodeid_basic():
    assert sanitize_nodeid("tests/foo.py::test_bar") == "tests_foo_py_test_bar"


def test_sanitize_nodeid_parametrized():
    s = sanitize_nodeid("tests/foo.py::test_bar[1-2-special!chars]")
    assert "test_bar" in s
    assert "!" not in s
    assert "[" not in s


def test_make_test_id_format():
    tid = make_test_id(7, "tests/x.py::test_y")
    assert tid.startswith("t007_")


def test_ms_between_iso_basic():
    a = "2026-05-22T12:00:00.000Z"
    b = "2026-05-22T12:00:01.500Z"
    assert ms_between_iso(a, b) == 1500


def test_validation_rejects_bad_status():
    from qa_dashboard.bundle.builder import build_test_record

    rec = build_test_record(
        test_id="t001_x",
        nodeid="x::y",
        name="y",
        file="x",
        line=1,
        status="banana",
        started_at="2026-01-01T00:00:00.000Z",
        ended_at="2026-01-01T00:00:01.000Z",
        duration_ms=1000,
        tags=[],
        parameters=None,
        outcome={
            "phase": "call",
            "longrepr": None,
            "traceback_summary": None,
            "stdout_tail": "",
            "stderr_tail": "",
        },
        step_data={"steps": [], "screenshots": []},
        run_dir=Path("."),
    )
    with pytest.raises(BundleValidationError):
        validate_test_record(rec)


def test_generates_index(sample_run):
    out = generate_index_html(sample_run)
    html = out.read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in html
    assert "const MANIFEST" in html
    assert "<style>" in html
    assert len(html) < 2_000_000
    # No unresolved template placeholders
    for ph in (
        "$title",
        "$styles",
        "$manifest_json",
        "$runtime_js",
        "$brand_style",
        "$brand_color_js",
    ):
        assert ph not in html, f"unresolved {ph}"


def test_generates_with_brand_color(sample_run):
    out = generate_index_html(sample_run, brand_color="#1d4ed8")
    html = out.read_text(encoding="utf-8")
    assert "--brand: #1d4ed8" in html


def test_rejects_invalid_brand_color(sample_run):
    # safe_color drops non-hex values; the rest of the page still renders
    out = generate_index_html(sample_run, brand_color="javascript:alert(1)")
    html = out.read_text(encoding="utf-8")
    assert "javascript:alert" not in html


def test_manifest_validates(sample_run):
    validate_manifest(json.loads((sample_run / "manifest.json").read_text()))


def test_all_test_records_validate(sample_run):
    for p in (sample_run / "tests").glob("*.json"):
        validate_test_record(json.loads(p.read_text()))


def test_standalone_html(sample_run):
    manifest = json.loads((sample_run / "manifest.json").read_text())
    test_id = manifest["tests"][0]["test_id"]
    out = generate_standalone_test_html(sample_run, test_id)
    html = out.read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in html
    # Screenshots should be inlined as data URIs
    if "data:image/png;base64," not in html:
        # only inlined if the test had screenshots
        rec = json.loads((sample_run / "tests" / f"{test_id}.json").read_text())
        assert not rec["screenshots"], "should have inlined screenshots"


def test_html_size_under_budget(sample_run):
    out = generate_index_html(sample_run)
    # 8 tests, budget is comfortably < 200 KB for the index page itself
    assert out.stat().st_size < 200_000


def test_handles_empty_run(tmp_path):
    """Run dir with zero tests should still produce a valid manifest+HTML."""
    run = tmp_path / "run_empty"
    (run / "tests").mkdir(parents=True)
    (run / "assets" / "screenshots").mkdir(parents=True)
    manifest = {
        "schema_version": SCHEMA_VERSION,
        "run_id": "run_empty",
        "started_at": "2026-01-01T00:00:00.000Z",
        "ended_at": "2026-01-01T00:00:01.000Z",
        "duration_ms": 1000,
        "status": "pass",
        "summary": {
            "total": 0,
            "passed": 0,
            "failed": 0,
            "skipped": 0,
            "error": 0,
            "xfailed": 0,
            "xpassed": 0,
            "pass_rate": 0.0,
        },
        "environment": {
            "pytest": "8.2.0",
            "python": "3.12",
            "platform": "test",
            "mitmproxy": "11.0",
            "qa_dashboard": "0.1.0",
            "hostname": "h",
            "user": "u",
            "cwd": "/",
        },
        "source": {
            "repo": None,
            "branch": None,
            "commit": None,
            "commit_message": None,
            "trigger": "local",
        },
        "config": {
            "mitm_port": 8080,
            "host_filter": "",
            "screenshot_on_fail": True,
            "command_line": "test",
        },
        "tests": [],
    }
    (run / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    out = generate_index_html(run)
    html = out.read_text()
    assert "<!DOCTYPE html>" in html


def test_plugin_runs_with_no_tests(tmp_path):
    """Plugin should produce a run directory even when 0 tests collected."""
    import subprocess

    p = tmp_path / "proj"
    p.mkdir()
    (p / "empty_test.py").write_text("# no tests\n")
    subprocess.run(
        ["pytest", str(p), "--qa-dashboard-mitm-skip"],
        capture_output=True,
        text=True,
        cwd=str(p),
        check=False,
    )
    # Pytest exits 5 when no tests collected. That's not an error for us.
    runs = list((p / "captures").glob("run_*"))
    assert len(runs) == 1
    m = json.loads((runs[0] / "manifest.json").read_text())
    assert m["summary"]["total"] == 0


def test_step_fixture_records_steps(tmp_path):
    import subprocess

    p = tmp_path / "proj"
    p.mkdir()
    (p / "test_steps.py").write_text(
        "def test_with_steps(step):\n"
        "    with step('one'):\n"
        "        pass\n"
        "    with step('two'):\n"
        "        step.screenshot(b'\\x89PNG\\r\\n\\x1a\\n' + b'\\x00'*100)\n"
    )
    subprocess.run(
        ["pytest", str(p), "--qa-dashboard-mitm-skip", "--qa-dashboard-no-report"],
        capture_output=True,
        text=True,
        cwd=str(p),
    )
    runs = list((p / "captures").glob("run_*"))
    assert len(runs) == 1
    tests = list((runs[0] / "tests").glob("*.json"))
    assert len(tests) == 1
    rec = json.loads(tests[0].read_text())
    assert rec["status"] == "pass"
    assert len(rec["steps"]) == 2
    assert rec["steps"][0]["name"] == "one"
    assert rec["steps"][1]["name"] == "two"
    assert len(rec["screenshots"]) == 1
    assert rec["screenshots"][0]["step_id"] == 2


def test_skipped_test_appears_in_manifest(tmp_path):
    """pytest.mark.skip short-circuits hooks; we must still record it."""
    import subprocess

    p = tmp_path / "proj"
    p.mkdir()
    (p / "test_s.py").write_text(
        "import pytest\n"
        "@pytest.mark.skip(reason='not yet')\n"
        "def test_a(): pass\n"
        "def test_b(): pass\n"
    )
    subprocess.run(
        ["pytest", str(p), "--qa-dashboard-mitm-skip", "--qa-dashboard-no-report"],
        capture_output=True,
        text=True,
        cwd=str(p),
    )
    runs = list((p / "captures").glob("run_*"))
    assert len(runs) == 1
    m = json.loads((runs[0] / "manifest.json").read_text())
    statuses = [t["status"] for t in m["tests"]]
    assert "skip" in statuses
    assert "pass" in statuses
    skipped = next(t for t in m["tests"] if t["status"] == "skip")
    assert "not yet" in (skipped.get("error_summary") or "")


def test_export_csv(sample_run, tmp_path):
    out = export_mod.export_csv(sample_run, tmp_path / "out.csv")
    assert out.exists()
    text = out.read_text(encoding="utf-8-sig")
    lines = text.strip().splitlines()
    assert lines[0].startswith("test_id,")
    # header + at least one row
    assert len(lines) >= 2


def test_export_json(sample_run, tmp_path):
    out = export_mod.export_json(sample_run, tmp_path / "out.json")
    assert out.exists()
    bundle = json.loads(out.read_text())
    assert "manifest" in bundle and "tests" in bundle
    assert len(bundle["tests"]) == bundle["manifest"]["summary"]["total"]


def test_export_zip(sample_run, tmp_path):
    out = export_mod.export_zip(sample_run, tmp_path / "out.zip")
    assert out.exists()
    import zipfile as _zip

    with _zip.ZipFile(out) as z:
        names = z.namelist()
    assert any(n.endswith("manifest.json") for n in names)
    assert any("tests/" in n and n.endswith(".json") for n in names)
    # Intermediate dirs should be excluded
    assert not any(n.startswith("_flows/") or n.startswith("_logs/") for n in names)


def test_export_xlsx(sample_run, tmp_path):
    pytest.importorskip("openpyxl")
    out = export_mod.export_xlsx(sample_run, tmp_path / "out.xlsx")
    assert out.exists()
    from openpyxl import load_workbook

    wb = load_workbook(out, read_only=True)
    assert set(wb.sheetnames) >= {"Summary", "Tests", "Steps", "Network"}


def test_export_docx(sample_run, tmp_path):
    pytest.importorskip("docx")
    out = export_mod.export_docx(sample_run, tmp_path / "out.docx")
    assert out.exists()
    from docx import Document

    doc = Document(str(out))
    assert any("Test Run Report" in p.text for p in doc.paragraphs) or any(
        "Test Run Report" in h.text for h in doc.paragraphs if h.style.name.startswith("Heading")
    )


def test_export_dispatcher_rejects_unknown(sample_run, tmp_path):
    with pytest.raises(ValueError):
        export_mod.export(sample_run, "rtf", tmp_path / "x.rtf")


def test_auto_screenshot_on_failure_without_step_fixture(tmp_path):
    """If a test has a driver-like fixture but doesn't use `step`, the
    plugin should still auto-capture a screenshot on failure."""
    import subprocess

    p = tmp_path / "proj"
    p.mkdir()
    # A "driver" fixture that mimics Selenium's API
    (p / "conftest.py").write_text(
        "import pytest\n"
        "class FakeDriver:\n"
        "    def get_screenshot_as_png(self):\n"
        "        return b'\\x89PNG\\r\\n\\x1a\\n' + b'\\x00'*200\n"
        "@pytest.fixture\n"
        "def driver():\n"
        "    return FakeDriver()\n"
    )
    (p / "test_x.py").write_text("def test_boom(driver):\n    assert False, 'kaboom'\n")
    subprocess.run(
        ["pytest", str(p), "--qa-dashboard-mitm-skip", "--qa-dashboard-no-report"],
        capture_output=True,
        text=True,
        cwd=str(p),
    )
    runs = list((p / "captures").glob("run_*"))
    assert len(runs) == 1
    rec = json.loads(list((runs[0] / "tests").glob("*.json"))[0].read_text())
    assert rec["status"] == "fail"
    assert len(rec["screenshots"]) == 1, "auto-screenshot should have been captured"
    assert "auto" in rec["screenshots"][0]["id"]
    # The synthetic step should be present
    assert len(rec["steps"]) == 1
    assert rec["steps"][0]["name"] == "Test execution"


def test_auto_screenshot_via_playwright_page_api(tmp_path):
    """Playwright uses ``page.screenshot()`` (no get_screenshot_as_png)."""
    import subprocess

    p = tmp_path / "proj"
    p.mkdir()
    (p / "conftest.py").write_text(
        "import pytest\n"
        "class FakePage:\n"
        "    def screenshot(self):\n"
        "        return b'\\x89PNG\\r\\n\\x1a\\n' + b'\\x00'*200\n"
        "@pytest.fixture\n"
        "def page():\n"
        "    return FakePage()\n"
    )
    (p / "test_x.py").write_text("def test_boom(page):\n    assert False\n")
    subprocess.run(
        ["pytest", str(p), "--qa-dashboard-mitm-skip", "--qa-dashboard-no-report"],
        capture_output=True,
        text=True,
        cwd=str(p),
    )
    runs = list((p / "captures").glob("run_*"))
    rec = json.loads(list((runs[0] / "tests").glob("*.json"))[0].read_text())
    assert len(rec["screenshots"]) == 1


def test_no_auto_screenshot_when_no_driver(tmp_path):
    """Tests with no driver-like fixture must not crash or fabricate."""
    import subprocess

    p = tmp_path / "proj"
    p.mkdir()
    (p / "test_x.py").write_text("def test_no_driver():\n    assert False\n")
    subprocess.run(
        ["pytest", str(p), "--qa-dashboard-mitm-skip", "--qa-dashboard-no-report"],
        capture_output=True,
        text=True,
        cwd=str(p),
    )
    runs = list((p / "captures").glob("run_*"))
    rec = json.loads(list((runs[0] / "tests").glob("*.json"))[0].read_text())
    assert rec["status"] == "fail"
    assert len(rec["screenshots"]) == 0


def test_export_via_cli(sample_run, tmp_path):
    from qa_dashboard.cli import main

    out = tmp_path / "out.csv"
    rc = main([str(sample_run), "--export", "csv", "-o", str(out)])
    assert rc == 0
    assert out.exists()


JUNIT_SUREFIRE_XML = """<?xml version="1.0" encoding="UTF-8"?>
<testsuite name="com.example.LoginTest" tests="4" failures="1" errors="1" skipped="1" time="3.5" timestamp="2026-05-23T11:00:35">
  <testcase classname="com.example.LoginTest" name="testValidLogin" time="1.234"/>
  <testcase classname="com.example.LoginTest" name="testInvalidPassword" time="0.567">
    <failure message="expected:&lt;true&gt; but was:&lt;false&gt;" type="AssertionError">java.lang.AssertionError: expected:&lt;true&gt; but was:&lt;false&gt;
\tat com.example.LoginTest.testInvalidPassword(LoginTest.java:42)</failure>
  </testcase>
  <testcase classname="com.example.LoginTest" name="testDatabaseDown" time="0.892">
    <error message="Connection refused" type="java.net.ConnectException">java.net.ConnectException: Connection refused</error>
  </testcase>
  <testcase classname="com.example.LoginTest" name="testFutureFeature" time="0.0">
    <skipped message="Not yet implemented"/>
  </testcase>
  <system-out>2026-05-23T11:00:35 INFO Starting LoginTest suite
2026-05-23T11:00:36 WARN Database slow response
2026-05-23T11:00:37 ERROR Failed to connect</system-out>
</testsuite>
"""


def test_ingest_junit_xml_basic(tmp_path):
    from qa_dashboard.bundle.schemas import validate_manifest, validate_test_record
    from qa_dashboard.ingest.junit_xml import ingest_junit_xml

    xml = tmp_path / "TEST-LoginTest.xml"
    xml.write_text(JUNIT_SUREFIRE_XML, encoding="utf-8")
    run = ingest_junit_xml([xml], out_root=tmp_path / "captures")
    manifest = json.loads((run / "manifest.json").read_text())
    assert manifest["summary"]["total"] == 4
    assert manifest["summary"]["passed"] == 1
    assert manifest["summary"]["failed"] == 1
    assert manifest["summary"]["error"] == 1
    assert manifest["summary"]["skipped"] == 1
    validate_manifest(manifest)
    for p in (run / "tests").glob("*.json"):
        validate_test_record(json.loads(p.read_text()))


def test_ingest_junit_xml_directory(tmp_path):
    from qa_dashboard.ingest.junit_xml import ingest_junit_xml

    xml_dir = tmp_path / "surefire-reports"
    xml_dir.mkdir()
    (xml_dir / "TEST-A.xml").write_text(JUNIT_SUREFIRE_XML, encoding="utf-8")
    (xml_dir / "TEST-B.xml").write_text(
        JUNIT_SUREFIRE_XML.replace("LoginTest", "PaymentTest"),
        encoding="utf-8",
    )
    run = ingest_junit_xml([xml_dir], out_root=tmp_path / "captures")
    m = json.loads((run / "manifest.json").read_text())
    assert m["summary"]["total"] == 8


def test_ingest_with_screenshots(tmp_path):
    from qa_dashboard.ingest.junit_xml import ingest_junit_xml

    xml = tmp_path / "TEST-x.xml"
    xml.write_text(JUNIT_SUREFIRE_XML, encoding="utf-8")
    shots = tmp_path / "screens"
    shots.mkdir()
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 200
    (shots / "testValidLogin_after.png").write_bytes(png)
    (shots / "testInvalidPassword_assert.png").write_bytes(png)
    run = ingest_junit_xml([xml], out_root=tmp_path / "captures", screenshots_dir=shots)
    valid = next(p for p in (run / "tests").glob("*testValidLogin*"))
    rec = json.loads(valid.read_text())
    assert len(rec["screenshots"]) == 1


def test_ingest_cli(tmp_path):
    from qa_dashboard.cli import main

    xml = tmp_path / "j.xml"
    xml.write_text(JUNIT_SUREFIRE_XML, encoding="utf-8")
    rc = main(["ingest", str(xml), "-o", str(tmp_path / "captures")])
    assert rc == 0
    runs = list((tmp_path / "captures").glob("run_*"))
    assert len(runs) == 1
    assert (runs[0] / "index.html").exists()
    assert (runs[0] / "manifest.json").exists()


def test_ingest_recursive_screenshot_match(tmp_path):
    """v0.5: screenshots in nested per-test dirs (Playwright JS style) auto-match."""
    from qa_dashboard.ingest.junit_xml import ingest_junit_xml

    xml = tmp_path / "TEST-x.xml"
    xml.write_text(JUNIT_SUREFIRE_XML, encoding="utf-8")
    shots = tmp_path / "test-results"
    # Playwright JS style: per-test subdir, generic png name
    (shots / "tests-testValidLogin-chromium").mkdir(parents=True)
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 200
    (shots / "tests-testValidLogin-chromium" / "test-failed-1.png").write_bytes(png)
    (shots / "tests-testInvalidPassword-chromium").mkdir(parents=True)
    (shots / "tests-testInvalidPassword-chromium" / "test-failed-1.png").write_bytes(png)
    run = ingest_junit_xml([xml], out_root=tmp_path / "captures", screenshots_dir=shots)
    valid = next(p for p in (run / "tests").glob("*testValidLogin*"))
    rec = json.loads(valid.read_text())
    assert len(rec["screenshots"]) == 1, "should pick up nested PNG by parent dir name"


def test_scaffold_playwright(tmp_path):
    """v0.5: `qa-dashboard scaffold playwright` copies the JS reporter."""
    from qa_dashboard.cli import main

    out = tmp_path / "qa-dashboard-reporter.js"
    rc = main(["scaffold", "playwright", "-o", str(out)])
    assert rc == 0
    assert out.exists()
    js = out.read_text(encoding="utf-8")
    # Sanity checks on the bundled reporter
    assert "QaDashboardReporter" in js
    assert "module.exports = QaDashboardReporter" in js
    assert "schema_version" in js
    assert "test.step" in js


def test_scaffold_playwright_refuses_overwrite(tmp_path):
    from qa_dashboard.cli import main

    out = tmp_path / "existing.js"
    out.write_text("// pre-existing\n", encoding="utf-8")
    rc = main(["scaffold", "playwright", "-o", str(out)])
    assert rc == 2  # refused
    assert out.read_text() == "// pre-existing\n"

    rc2 = main(["scaffold", "playwright", "-o", str(out), "--force"])
    assert rc2 == 0
    assert "QaDashboardReporter" in out.read_text()


def test_playwright_reporter_js_is_packaged():
    """The shipped wheel must include the JS file."""
    from qa_dashboard.integrations import playwright as pw_pkg

    js_path = Path(pw_pkg.__file__).parent / "qa-dashboard-reporter.js"
    assert js_path.exists(), f"missing reporter at {js_path}"
    assert js_path.stat().st_size > 5000, "reporter looks suspiciously small"


def test_failing_test_captures_outcome(tmp_path):
    import subprocess

    p = tmp_path / "proj"
    p.mkdir()
    (p / "test_fail.py").write_text("def test_boom():\n    assert False, 'kaboom'\n")
    subprocess.run(
        ["pytest", str(p), "--qa-dashboard-mitm-skip", "--qa-dashboard-no-report"],
        capture_output=True,
        text=True,
        cwd=str(p),
    )
    runs = list((p / "captures").glob("run_*"))
    rec = json.loads(list((runs[0] / "tests").glob("*.json"))[0].read_text())
    assert rec["status"] == "fail"
    assert "kaboom" in (rec["outcome"]["longrepr"] or "")
